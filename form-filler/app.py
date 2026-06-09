import hashlib
import io
import json
import logging
import os
import re
import time
from datetime import datetime, timezone
from json import JSONDecodeError
from pathlib import Path
from typing import Any, Dict, Mapping, Optional
from zoneinfo import ZoneInfo

import openpyxl
import requests
from fastapi import Body, File, Form, FastAPI, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from psycopg2 import connect, sql
from psycopg2.extras import RealDictCursor

app = FastAPI(title="n8n XLSX Form Filler", version="3.0.0")
logger = logging.getLogger("form-filler")
APP_DIR = Path(__file__).resolve().parent
REPORT_UI_DIST_DIR = APP_DIR / "report-ui" / "dist"
REPORT_UI_ASSETS_DIR = REPORT_UI_DIST_DIR / "assets"

app.mount("/assets", StaticFiles(directory=str(REPORT_UI_ASSETS_DIR), check_dir=False), name="report-ui-assets")


DEFAULT_SUPPORT_REQUEST_FIELD_MAP: Dict[str, str] = {
    "appointment_date": "B2",
    "appointment_time": "D2",
    "end_date": "B3",
    "end_time": "D3",
    "address": "B4",
    "google_map_link": "B5",
    "request_details": "B6",
    "current_condition": "B7",
    "equipment_prepared_by_tpt": "B8",
    "equipment_handover_plan": "D8",
    "issued_date": "B9",
    "issued_by": "D9",
    "contact_person": "B10",
    "contact_number": "D10",
    "sales_rep_participation": "C11",
    "customer_name": "C12",
    "sales_channel": "C13",
}

LANGUAGE_ORDER = ("en", "ja", "th")
LANGUAGE_NAMES = {
    "en": "English",
    "ja": "Japanese",
    "th": "Thai",
}
LANGUAGE_LABELS = {
    "en": "EN",
    "ja": "JA",
    "th": "TH",
}
LANGUAGE_PRIORITY = ("en", "ja", "th")
REPLY_META_SHEET_NAME = "_reply_meta"
REPLY_META_VERSION = "maintenance_reply_v1"
MAX_EXCEL_CELL_LENGTH = 32767

MAINTENANCE_NORMAL_FIELD_SPECS: Dict[str, Dict[str, Any]] = {
    "customer_appointment_date": {
        "cell": "B2",
        "kind": "appointment_part",
        "db_column": "appointment_start",
        "part": "date",
        "translate_to_english": False,
    },
    "customer_appointment_time": {
        "cell": "D2",
        "kind": "appointment_part",
        "db_column": "appointment_start",
        "part": "time",
        "translate_to_english": False,
    },
    "end_date": {
        "cell": "B3",
        "kind": "appointment_part",
        "db_column": "appointment_end",
        "part": "date",
        "translate_to_english": False,
    },
    "end_time": {
        "cell": "D3",
        "kind": "appointment_part",
        "db_column": "appointment_end",
        "part": "time",
        "translate_to_english": False,
    },
    "address": {
        "cell": "B4",
        "kind": "column",
        "db_column": "service_address",
        "translate_to_english": True,
    },
    "google_map_link": {
        "cell": "B5",
        "kind": "column",
        "db_column": "google_map_link",
        "translate_to_english": False,
    },
    "equipment_prepared_by_tpt": {
        "cell": "B8",
        "kind": "column",
        "db_column": "equipment_prepared_by_tpt",
        "translate_to_english": True,
    },
    "equipment_handover_plan": {
        "cell": "D8",
        "kind": "column",
        "db_column": "equipment_handover_plan",
        "translate_to_english": True,
    },
    "issued_date": {
        "cell": "B9",
        "kind": "column",
        "db_column": "issued_date",
        "translate_to_english": True,
    },
    "issued_by": {
        "cell": "D9",
        "kind": "column",
        "db_column": "issued_by",
        "translate_to_english": True,
    },
    "contact_person": {
        "cell": "B10",
        "kind": "column",
        "db_column": "contact_person",
        "translate_to_english": True,
    },
    "contact_number": {
        "cell": "D10",
        "kind": "column",
        "db_column": "contact_number",
        "translate_to_english": False,
    },
    "sales_representative_participation": {
        "cell": "C11",
        "kind": "column",
        "db_column": "sales_representative_participation",
        "translate_to_english": True,
    },
    "customer_name": {
        "cell": "C12",
        "kind": "column",
        "db_column": "customer_name",
        "translate_to_english": True,
    },
    "sales_channel": {
        "cell": "C13",
        "kind": "column",
        "db_column": "sales_channel",
        "translate_to_english": True,
    },
}

MAINTENANCE_MULTILINGUAL_FIELD_SPECS: Dict[str, Dict[str, Any]] = {
    "request_details": {
        "cell": "B6",
        "db_column": "request_details",
        "source_lang_column": "request_details_source_lang",
        "translation_columns": {
            "en": "request_details_en",
            "ja": "request_details_ja",
            "th": "request_details_th",
        },
        "label": "Request details",
    },
    "current_condition": {
        "cell": "B7",
        "db_column": "current_condition",
        "source_lang_column": "current_condition_source_lang",
        "translation_columns": {
            "en": "current_condition_en",
            "ja": "current_condition_ja",
            "th": "current_condition_th",
        },
        "label": "Current conditions",
    },
}

MAINTENANCE_SCHEMA_ADDITIONS = {
    "equipment_prepared_by_tpt": "TEXT",
    "equipment_handover_plan": "TEXT",
    "issued_date": "TEXT",
    "issued_by": "TEXT",
    "sales_representative_participation": "TEXT",
    "sales_channel": "TEXT",
    "request_details_source_lang": "VARCHAR(5)",
    "request_details_en": "TEXT",
    "request_details_ja": "TEXT",
    "request_details_th": "TEXT",
    "current_condition_source_lang": "VARCHAR(5)",
    "current_condition_en": "TEXT",
    "current_condition_ja": "TEXT",
    "current_condition_th": "TEXT",
}


def _maintenance_field_cells() -> Dict[str, str]:
    field_cells: Dict[str, str] = {}
    for key, spec in MAINTENANCE_NORMAL_FIELD_SPECS.items():
        field_cells[key] = spec["cell"]
    for key, spec in MAINTENANCE_MULTILINGUAL_FIELD_SPECS.items():
        field_cells[key] = spec["cell"]
    return field_cells


def _get_env(name: str, default: str = "") -> str:
    value = os.getenv(name)
    if value is None:
        return default
    value = value.strip()
    return value if value else default


def _split_csv_env(raw: str) -> list[str]:
    return [item for item in (_normalize_text(part) for part in raw.split(",")) if item]


def _normalize_key(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", value.strip().lower())
    return normalized.strip("_")


def _load_support_request_field_map() -> Dict[str, str]:
    raw = os.getenv("FIELD_MAP_JSON")
    if not raw:
        return DEFAULT_SUPPORT_REQUEST_FIELD_MAP
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Invalid FIELD_MAP_JSON: {exc}") from exc
    if not isinstance(data, dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in data.items()):
        raise RuntimeError("FIELD_MAP_JSON must be a JSON object of string keys and cell addresses")
    return data


def _safe_filename(filename: str) -> str:
    cleaned = filename.replace("\\", "_").replace("/", "_").strip()
    if not cleaned:
        return "output.xlsx"
    cleaned = re.sub(r"[<>:\"|?*]+", "_", cleaned)
    if not cleaned.lower().endswith(".xlsx"):
        cleaned += ".xlsx"
    return cleaned


def _get_display_timezone() -> ZoneInfo:
    timezone_name = _get_env("FORM_TIMEZONE", _get_env("TZ", "UTC"))
    try:
        return ZoneInfo(timezone_name)
    except Exception:
        return ZoneInfo("UTC")


def _ordinal(day: int) -> str:
    if 10 <= day % 100 <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
    return f"{day}{suffix}"


def _normalize_text(value: Any) -> str:
    text = _stringify(value)
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def _dedupe_strings_preserve_order(items: list[str]) -> list[str]:
    deduped: list[str] = []
    seen: set[str] = set()
    for item in items:
        key = item.lower()
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)
    return deduped


def _hash_text(value: Any) -> str:
    return hashlib.sha256(_normalize_text(value).encode("utf-8")).hexdigest()


def _coerce_datetime(value: Any) -> Optional[datetime]:
    if value is None:
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value).strip()
        if not text:
            return None
        try:
            dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            return None
    if dt.tzinfo is not None:
        return dt.astimezone(_get_display_timezone())
    return dt


def _format_date_for_form(value: Any) -> str:
    dt = _coerce_datetime(value)
    if dt is None:
        return str(value or "").strip()
    return f"{_ordinal(dt.day)} {dt.strftime('%b %Y')}"


def _format_time_for_form(value: Any) -> str:
    dt = _coerce_datetime(value)
    if dt is None:
        return str(value or "").strip()
    return dt.strftime("%I.%M%p")


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = _coerce_datetime(value) or value
        return value.isoformat(sep=" ", timespec="minutes")
    return str(value).strip()


def _extract_json_object(text: str) -> Dict[str, Any]:
    decoder = json.JSONDecoder()
    for index, char in enumerate(text):
        if char != "{":
            continue
        try:
            payload, _ = decoder.raw_decode(text[index:])
        except JSONDecodeError:
            continue
        if isinstance(payload, dict):
            return payload
    raise ValueError(f"Model output did not contain a JSON object: {text[:300]}")


def _looks_english(text: str) -> bool:
    compact = "".join(ch for ch in _normalize_text(text) if not ch.isspace())
    if not compact:
        return True
    ascii_count = sum(1 for ch in compact if ord(ch) < 128)
    return ascii_count / len(compact) >= 0.9


def _normalize_source_language(candidate: Any, source_text: str = "") -> str:
    normalized = _normalize_key(str(candidate or ""))
    aliases = {
        "en": "en",
        "english": "en",
        "ja": "ja",
        "jp": "ja",
        "japanese": "ja",
        "th": "th",
        "thai": "th",
    }
    if normalized in aliases:
        return aliases[normalized]
    if source_text and _looks_english(source_text):
        return "en"
    return "en"


def _get_ollama_think_setting() -> Any:
    raw = _get_env("OLLAMA_THINK", "false").lower()
    if raw in {"true", "1", "yes", "on"}:
        return True
    if raw in {"false", "0", "no", "off"}:
        return False
    return raw


def _ollama_chat(
    messages: list[dict[str, str]],
    expect_json: bool = True,
    model: str | None = None,
) -> str:
    base_url = _get_env("OLLAMA_BASE_URL", "http://ollama:11434").rstrip("/")
    model = model or _get_env("OLLAMA_MODEL", "kamekichi128/qwen3-4b-instruct-2507:latest")
    timeout_seconds = int(_get_env("OLLAMA_TIMEOUT_SECONDS", "180"))
    payload: Dict[str, Any] = {
        "model": model,
        "messages": messages,
        "stream": False,
        "think": _get_ollama_think_setting(),
        "options": {"temperature": 0},
    }
    if expect_json:
        payload["format"] = "json"
    try:
        response = requests.post(
            f"{base_url}/api/chat",
            json=payload,
            timeout=timeout_seconds,
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        raise RuntimeError(f"Failed to call Ollama at {base_url}: {exc}") from exc
    data = response.json()
    message = data.get("message") or {}
    content = message.get("content")
    if not isinstance(content, str) or not content.strip():
        thinking = message.get("thinking")
        if isinstance(thinking, str) and thinking.strip():
            raise RuntimeError(
                f"Ollama returned only a thinking trace for model {model}. "
                "Set OLLAMA_THINK=false or use a non-thinking model."
            )
        raise RuntimeError(f"Ollama returned an empty response: {data}")
    return content.strip()


def _deepseek_chat(messages: list[dict[str, str]], expect_json: bool = True) -> str:
    api_key = _get_env("DEEPSEEK_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("DEEPSEEK_API_KEY is not set.")

    base_url = _get_env("DEEPSEEK_BASE_URL", "https://api.deepseek.com").rstrip("/")
    model = _get_env("DEEPSEEK_MODEL", "deepseek-chat")
    timeout_seconds = int(_get_env("DEEPSEEK_TIMEOUT_SECONDS", "180"))
    payload: Dict[str, Any] = {
        "model": model,
        "messages": messages,
        "stream": False,
        "temperature": 0,
        "thinking": {"type": "disabled"},
    }
    if expect_json:
        payload["response_format"] = {"type": "json_object"}

    try:
        response = requests.post(
            f"{base_url}/chat/completions",
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
            },
            json=payload,
            timeout=timeout_seconds,
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        detail = ""
        if exc.response is not None:
            try:
                detail = exc.response.text
            except Exception:
                detail = ""
        if detail:
            raise RuntimeError(f"Failed to call DeepSeek at {base_url}: {exc}; response={detail}") from exc
        raise RuntimeError(f"Failed to call DeepSeek at {base_url}: {exc}") from exc

    data = response.json()
    choices = data.get("choices") or []
    if not choices:
        raise RuntimeError(f"DeepSeek returned no choices: {data}")
    message = choices[0].get("message") or {}
    content = message.get("content")
    if not isinstance(content, str) or not content.strip():
        raise RuntimeError(f"DeepSeek returned an empty response: {data}")
    return content.strip()


def _translate_fields_to_english(values: Mapping[str, str]) -> Dict[str, str]:
    result: Dict[str, str] = {}
    payload: Dict[str, str] = {}
    for key, value in values.items():
        if not _normalize_text(value):
            continue
        if _looks_english(value):
            result[key] = _normalize_text(value)
        else:
            payload[key] = _normalize_text(value)
    if not payload:
        return result
    content = _ollama_chat(
        [
            {
                "role": "system",
                "content": (
                    "You translate maintenance form values into natural English. "
                    "Return JSON only. Keep the same keys. Keep URLs, phone numbers, "
                    "names, model numbers, and dates accurate."
                ),
            },
            {
                "role": "user",
                "content": (
                    "Translate every value in the JSON object to English. "
                    "Return only a JSON object with exactly the same keys.\n"
                    f"{json.dumps(payload, ensure_ascii=False)}"
                ),
            },
        ]
    )
    translated = _extract_json_object(content)
    for key, original in payload.items():
        translated_value = translated.get(key, original)
        result[key] = _normalize_text(translated_value) or original
    return result


def _translate_text_to_language(text: str, target_language: str, field_name: str) -> str:
    source = _normalize_text(text)
    if not source:
        return ""
    content = _ollama_chat(
        [
            {
                "role": "system",
                "content": (
                    "You translate maintenance form text. "
                    "Return only the translated text with no JSON, no labels, and no commentary."
                ),
            },
            {
                "role": "user",
                "content": (
                    f"Translate the following {field_name} text into {target_language}. "
                    "Preserve technical meaning, product names, numbers, and paragraph breaks. "
                    "Return only the translation.\n\n"
                    f"{source}"
                ),
            },
        ],
        expect_json=False,
    )
    return _normalize_text(content)


def _build_language_triplet(source_text: str, source_lang: str, field_name: str) -> Dict[str, str]:
    canonical = _normalize_text(source_text)
    if not canonical:
        return {lang: "" for lang in LANGUAGE_ORDER}
    resolved_source_lang = _normalize_source_language(source_lang, canonical)
    blocks: Dict[str, str] = {}
    for lang in LANGUAGE_ORDER:
        if lang == resolved_source_lang:
            blocks[lang] = canonical
        elif lang == "en" and resolved_source_lang == "en":
            blocks[lang] = canonical
        else:
            blocks[lang] = _translate_text_to_language(canonical, LANGUAGE_NAMES[lang], field_name)
    return blocks


def _format_labeled_blocks(blocks: Mapping[str, str]) -> str:
    parts = []
    for lang in LANGUAGE_ORDER:
        label = LANGUAGE_LABELS[lang]
        parts.append(f"[{label}]\n{_normalize_text(blocks.get(lang, ''))}")
    return "\n\n".join(parts)


def _parse_labeled_blocks(value: str) -> Dict[str, str]:
    normalized = _normalize_text(value)
    pattern = re.compile(r"(?m)^\[(EN|JA|TH)\]\s*$")
    matches = list(pattern.finditer(normalized))
    if [match.group(1) for match in matches] != ["EN", "JA", "TH"]:
        raise ValueError("Expected labeled blocks [EN], [JA], [TH]")
    blocks: Dict[str, str] = {}
    for index, match in enumerate(matches):
        start = match.end()
        if start < len(normalized) and normalized[start] == "\n":
            start += 1
        end = matches[index + 1].start() if index + 1 < len(matches) else len(normalized)
        blocks[match.group(1).lower()] = normalized[start:end].strip()
    return blocks


def _connect_postgres():
    host = _get_env("POSTGRES_HOST", "host.docker.internal")
    port = int(_get_env("POSTGRES_PORT", "5432"))
    database = _get_env("POSTGRES_DATABASE")
    user = _get_env("POSTGRES_USER")
    password = _get_env("POSTGRES_PASSWORD")
    if not database or not user or not password:
        raise RuntimeError(
            "POSTGRES_DATABASE, POSTGRES_USER, and POSTGRES_PASSWORD must be configured for maintenance form generation"
        )
    return connect(
        host=host,
        port=port,
        dbname=database,
        user=user,
        password=password,
        cursor_factory=RealDictCursor,
    )


def _resolve_table_name(candidate: Any = None) -> str:
    resolved = _normalize_text(candidate) or _get_env("POSTGRES_TABLE", "n8n_info_record")
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*([.][A-Za-z_][A-Za-z0-9_]*)?", resolved):
        raise HTTPException(status_code=422, detail="table_name is invalid")
    return resolved


def _postgres_table_identifier(table_name: Optional[str] = None):
    return sql.Identifier(*_resolve_table_name(table_name).split("."))


def _ensure_maintenance_schema(connection, table_name: Optional[str] = None) -> None:
    table_identifier = _postgres_table_identifier(table_name)
    with connection.cursor() as cursor:
        for column_name, column_type in MAINTENANCE_SCHEMA_ADDITIONS.items():
            cursor.execute(
                sql.SQL("ALTER TABLE {} ADD COLUMN IF NOT EXISTS {} {}").format(
                    table_identifier,
                    sql.Identifier(column_name),
                    sql.SQL(column_type),
                )
            )


def _fetch_record(connection, record_id: Optional[int], table_name: Optional[str] = None) -> Dict[str, Any]:
    table_identifier = _postgres_table_identifier(table_name)
    with connection.cursor() as cursor:
        if record_id is None:
            query = sql.SQL("SELECT * FROM {} ORDER BY id DESC LIMIT 1").format(table_identifier)
            cursor.execute(query)
        else:
            query = sql.SQL("SELECT * FROM {} WHERE id = %s LIMIT 1").format(table_identifier)
            cursor.execute(query, (record_id,))
        row = cursor.fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="No matching maintenance record found")
    return dict(row)


def _update_record_columns(
    connection,
    record_id: int,
    updates: Mapping[str, Any],
    table_name: Optional[str] = None,
) -> None:
    if not updates:
        return
    table_identifier = _postgres_table_identifier(table_name)
    assignments = []
    values = []
    for column_name, value in updates.items():
        assignments.append(sql.SQL("{} = %s").format(sql.Identifier(column_name)))
        values.append(value)
    values.append(record_id)
    query = sql.SQL("UPDATE {} SET {} WHERE id = %s").format(
        table_identifier,
        sql.SQL(", ").join(assignments),
    )
    with connection.cursor() as cursor:
        cursor.execute(query, values)
        if cursor.rowcount != 1:
            raise HTTPException(status_code=404, detail=f"Record {record_id} not found for update")


def _build_normal_export_values(record: Mapping[str, Any]) -> Dict[str, str]:
    raw_values: Dict[str, str] = {}
    for key, spec in MAINTENANCE_NORMAL_FIELD_SPECS.items():
        if spec["kind"] == "appointment_part":
            source_datetime = record.get(spec["db_column"])
            if spec["part"] == "date":
                raw_values[key] = _format_date_for_form(source_datetime)
            else:
                raw_values[key] = _format_time_for_form(source_datetime)
            continue
        raw_values[key] = _normalize_text(record.get(spec["db_column"]))

    english_targets = {
        key: value
        for key, value in raw_values.items()
        if value and MAINTENANCE_NORMAL_FIELD_SPECS[key]["translate_to_english"]
    }
    translated = _translate_fields_to_english(english_targets)
    final_values: Dict[str, str] = {}
    for key, value in raw_values.items():
        final_values[key] = translated.get(key, value)
    return final_values


def _prepare_multilingual_export_value(record: Mapping[str, Any], field_key: str) -> Dict[str, Any]:
    spec = MAINTENANCE_MULTILINGUAL_FIELD_SPECS[field_key]
    db_column = spec["db_column"]
    source_lang_column = spec["source_lang_column"]
    translation_columns = spec["translation_columns"]
    canonical = _normalize_text(record.get(db_column))
    source_lang = _normalize_source_language(record.get(source_lang_column) or record.get("language"), canonical)
    stored_blocks = {
        lang: _normalize_text(record.get(column_name))
        for lang, column_name in translation_columns.items()
    }
    updates: Dict[str, Any] = {}

    if not canonical and any(stored_blocks.values()):
        preferred_langs = [source_lang] + [lang for lang in LANGUAGE_PRIORITY if lang != source_lang]
        selected_lang = next((lang for lang in preferred_langs if stored_blocks.get(lang)), "en")
        canonical = stored_blocks.get(selected_lang, "")
        source_lang = selected_lang
        updates[db_column] = canonical
        updates[source_lang_column] = selected_lang

    if not canonical:
        blocks = {lang: "" for lang in LANGUAGE_ORDER}
    else:
        needs_regeneration = any(not stored_blocks.get(lang) for lang in LANGUAGE_ORDER)
        if stored_blocks.get(source_lang) and stored_blocks[source_lang] != canonical:
            needs_regeneration = True
        if needs_regeneration:
            blocks = _build_language_triplet(canonical, source_lang, spec["label"])
        else:
            blocks = dict(stored_blocks)
            blocks[source_lang] = canonical

        for lang, column_name in translation_columns.items():
            if record.get(column_name) != blocks[lang]:
                updates[column_name] = blocks[lang]
        if record.get(source_lang_column) != source_lang:
            updates[source_lang_column] = source_lang
        if record.get(db_column) != canonical:
            updates[db_column] = canonical

    return {
        "field_key": field_key,
        "db_column": db_column,
        "cell": spec["cell"],
        "source_lang": source_lang,
        "source_text": canonical,
        "blocks": blocks,
        "display_text": _format_labeled_blocks(blocks),
        "db_updates": updates,
    }


def _metadata_payload(
    record_id: int,
    table_name: str,
    sheet_name: str,
    normal_values: Mapping[str, str],
    multilingual_values: Mapping[str, Mapping[str, Any]],
) -> str:
    metadata = {
        "version": REPLY_META_VERSION,
        "record_id": record_id,
        "table_name": table_name,
        "sheet_name": sheet_name,
        "field_cells": _maintenance_field_cells(),
        "normal_fields": {
            field_key: {
                "cell": MAINTENANCE_NORMAL_FIELD_SPECS[field_key]["cell"],
                "original": value,
            }
            for field_key, value in normal_values.items()
        },
        "multilingual_fields": {
            field_key: {
                "cell": payload["cell"],
                "source_lang": payload["source_lang"],
                "source_text": payload["source_text"],
                "blocks": payload["blocks"],
                "hashes": {lang: _hash_text(text) for lang, text in payload["blocks"].items()},
            }
            for field_key, payload in multilingual_values.items()
        },
    }
    serialized = json.dumps(metadata, ensure_ascii=False, separators=(",", ":"))
    if len(serialized) > MAX_EXCEL_CELL_LENGTH:
        raise HTTPException(status_code=500, detail="Reply metadata is too large to fit into the workbook")
    return serialized


def _upsert_reply_metadata_sheet(workbook: openpyxl.Workbook, metadata_text: str) -> None:
    if REPLY_META_SHEET_NAME in workbook.sheetnames:
        metadata_sheet = workbook[REPLY_META_SHEET_NAME]
        workbook.remove(metadata_sheet)
    metadata_sheet = workbook.create_sheet(REPLY_META_SHEET_NAME)
    metadata_sheet["A1"] = metadata_text
    metadata_sheet.sheet_state = "hidden"


def _render_workbook_to_bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _save_output_copy(content: bytes, filename: str) -> Optional[str]:
    output_dir = _get_env("OUTPUT_DIR", "/outputs")
    if not output_dir:
        return None
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, filename)
    with open(output_path, "wb") as handle:
        handle.write(content)
    return output_path


def _stream_xlsx(
    content: bytes,
    filename: str,
    output_path: Optional[str] = None,
    extra_headers: Optional[Mapping[str, str]] = None,
) -> StreamingResponse:
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    if output_path:
        headers["X-Output-Path"] = output_path
    if extra_headers:
        headers.update({key: value for key, value in extra_headers.items() if value})
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


def _load_reply_metadata(workbook: openpyxl.Workbook) -> Dict[str, Any]:
    if REPLY_META_SHEET_NAME not in workbook.sheetnames:
        raise HTTPException(status_code=422, detail="Workbook is missing reply metadata")
    metadata_text = workbook[REPLY_META_SHEET_NAME]["A1"].value
    if not isinstance(metadata_text, str) or not metadata_text.strip():
        raise HTTPException(status_code=422, detail="Reply metadata is empty")
    try:
        metadata = json.loads(metadata_text)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=422, detail="Reply metadata is invalid JSON") from exc
    if metadata.get("version") != REPLY_META_VERSION:
        raise HTTPException(status_code=422, detail="Workbook reply metadata version is unsupported")
    return metadata


def _parse_form_date(value: str) -> datetime.date:
    normalized = _normalize_text(value)
    if not normalized:
        raise ValueError("date is empty")
    normalized = re.sub(r"(?<=\d)(st|nd|rd|th)\b", "", normalized, flags=re.IGNORECASE)
    formats = ("%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
    for fmt in formats:
        try:
            return datetime.strptime(normalized, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Unsupported date format: {value}")


def _parse_form_time(value: str) -> datetime.time:
    normalized = _normalize_text(value).upper().replace(" ", "")
    if not normalized:
        raise ValueError("time is empty")
    formats = ("%I.%M%p", "%I:%M%p", "%H:%M", "%H.%M")
    for fmt in formats:
        try:
            return datetime.strptime(normalized, fmt).time()
        except ValueError:
            continue
    raise ValueError(f"Unsupported time format: {value}")


def _parse_form_timestamp(date_value: str, time_value: str) -> datetime:
    local_date = _parse_form_date(date_value)
    local_time = _parse_form_time(time_value)
    timezone_value = _get_display_timezone()
    aware = datetime.combine(local_date, local_time).replace(tzinfo=timezone_value)
    return aware.replace(tzinfo=None)


def _split_resolved_table_name(table_name: str) -> tuple[str, str]:
    parts = _resolve_table_name(table_name).split(".")
    if len(parts) == 1:
        return "public", parts[0]
    return parts[0], parts[1]


def _report_allowed_tables() -> list[str]:
    configured = _split_csv_env(_get_env("REPORT_ALLOWED_TABLES", "info_record,n8n_info_record"))
    if not configured:
        configured = ["info_record", "n8n_info_record"]
    return _dedupe_strings_preserve_order([_resolve_table_name(item) for item in configured])


def _report_default_table() -> str:
    allowed_tables = _report_allowed_tables()
    configured_default = _normalize_text(_get_env("REPORT_DEFAULT_TABLE"))
    if not configured_default:
        return allowed_tables[0]
    resolved_default = _resolve_table_name(configured_default)
    return resolved_default if resolved_default in allowed_tables else allowed_tables[0]


def _ensure_report_table_allowed(table_name: str) -> str:
    resolved = _resolve_table_name(table_name)
    if resolved not in _report_allowed_tables():
        raise HTTPException(status_code=422, detail="table_name is not allowed for the report UI")
    return resolved


def _table_label(table_name: str) -> str:
    leaf_name = table_name.split(".")[-1]
    label = re.sub(r"[_\s]+", " ", leaf_name).strip()
    return label.upper() if label.lower() == "n8n" else label.title()


def _utcnow_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _list_table_columns(connection, table_name: str) -> list[Dict[str, Any]]:
    schema_name, relation_name = _split_resolved_table_name(table_name)
    query = (
        "SELECT column_name, data_type, udt_name, is_nullable, ordinal_position "
        "FROM information_schema.columns "
        "WHERE table_schema = %s AND table_name = %s "
        "ORDER BY ordinal_position"
    )
    with connection.cursor() as cursor:
        cursor.execute(query, (schema_name, relation_name))
        rows = cursor.fetchall()
    if not rows:
        raise HTTPException(status_code=404, detail=f"Table not found: {table_name}")
    return [dict(row) for row in rows]


def _count_table_rows(connection, table_name: str) -> int:
    query = sql.SQL("SELECT COUNT(*) AS row_count FROM {}").format(_postgres_table_identifier(table_name))
    with connection.cursor() as cursor:
        cursor.execute(query)
        row = cursor.fetchone()
    return int((row or {}).get("row_count") or 0)


def _json_safe_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, datetime):
        normalized = _coerce_datetime(value) or value
        return normalized.isoformat(sep=" ", timespec="seconds")
    try:
        json.dumps(value, ensure_ascii=False)
        return value
    except TypeError:
        return str(value)


def _sample_table_rows(connection, table_name: str, columns: list[Dict[str, Any]], limit: int) -> list[Dict[str, Any]]:
    column_names = {column["column_name"] for column in columns}
    order_column = next(
        (name for name in ("updated_at", "created_at", "created_on", "appointment_start", "id") if name in column_names),
        None,
    )
    order_clause = sql.SQL("")
    if order_column:
        order_clause = sql.SQL(" ORDER BY {} DESC").format(sql.Identifier(order_column))
    query = sql.SQL("SELECT * FROM {}{} LIMIT %s").format(_postgres_table_identifier(table_name), order_clause)
    with connection.cursor() as cursor:
        cursor.execute(query, (limit,))
        rows = cursor.fetchall()
    return [{key: _json_safe_value(value) for key, value in dict(row).items()} for row in rows]


def _prepare_rows_for_llm(rows: list[Dict[str, Any]], max_chars: int = 12000) -> Dict[str, Any]:
    prepared_rows: list[Dict[str, Any]] = []
    total_chars = 0
    truncated = False
    for row in rows:
        prepared_row: Dict[str, Any] = {}
        for key, value in row.items():
            safe_value = _json_safe_value(value)
            if isinstance(safe_value, str) and len(safe_value) > 300:
                safe_value = safe_value[:300].rstrip() + "..."
            prepared_row[key] = safe_value
        serialized = json.dumps(prepared_row, ensure_ascii=False, separators=(",", ":"))
        if prepared_rows and total_chars + len(serialized) > max_chars:
            truncated = True
            break
        prepared_rows.append(prepared_row)
        total_chars += len(serialized)
    return {"rows": prepared_rows, "truncated": truncated}


def _normalize_string_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [_normalize_text(item) for item in value if _normalize_text(item)]
    if isinstance(value, str) and _normalize_text(value):
        return [_normalize_text(value)]
    return []


def _dedupe_strings(items: list[str], limit: int = 4) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for item in items:
        normalized = _normalize_text(item)
        if not normalized:
            continue
        key = normalized.lower()
        if key in seen:
            continue
        seen.add(key)
        result.append(normalized)
        if len(result) >= limit:
            break
    return result


def _matching_columns(column_names: list[str], keywords: tuple[str, ...]) -> list[str]:
    return [name for name in column_names if any(keyword in name for keyword in keywords)]


def _column_presence_stats(sample_rows: list[Dict[str, Any]]) -> Dict[str, Dict[str, int]]:
    stats: Dict[str, Dict[str, int]] = {}
    for row in sample_rows:
        for key, value in row.items():
            entry = stats.setdefault(key, {"nonempty": 0, "total": 0})
            entry["total"] += 1
            if _normalize_text(value):
                entry["nonempty"] += 1
    return stats


def _fallback_analysis(
    table_name: str,
    row_count: int,
    columns: list[Dict[str, Any]],
    sample_rows: list[Dict[str, Any]],
    sample_truncated: bool,
) -> Dict[str, Any]:
    column_names = [str(column.get("column_name") or "").lower() for column in columns]
    time_columns = _matching_columns(column_names, ("date", "time", "created", "updated", "appointment", "start", "end"))
    workflow_columns = _matching_columns(column_names, ("status", "stage", "approval", "review", "owner", "assignee", "department"))
    customer_columns = _matching_columns(column_names, ("customer", "contact", "phone", "email", "address", "sales", "channel"))
    request_columns = _matching_columns(column_names, ("request", "detail", "condition", "issue", "problem", "equipment", "service"))
    language_columns = [name for name in column_names if name.endswith(("_en", "_ja", "_th")) or "language" in name]
    completion_columns = _matching_columns(column_names, ("result", "resolution", "complete", "closed", "outcome"))
    priority_columns = _matching_columns(column_names, ("priority", "severity", "urgent"))
    feedback_columns = _matching_columns(column_names, ("feedback", "satisfaction", "rating", "score"))
    cost_columns = _matching_columns(column_names, ("cost", "price", "amount", "expense", "revenue"))

    stats = _column_presence_stats(sample_rows)
    sparse_columns = [
        name
        for name, stat in sorted(stats.items(), key=lambda item: (item[1]["nonempty"], item[0]))
        if stat["total"] > 0 and stat["nonempty"] < stat["total"]
    ][:3]

    default_status = "empty" if row_count == 0 else "insufficient" if row_count < 5 else "usable"
    if row_count == 0:
        executive_summary = (
            f"The `{table_name}` table has {len(columns)} columns but no rows yet. "
            "The priority should be stable data collection, mandatory field enforcement, and KPI definition before using the dataset for trend analysis."
        )
    elif row_count < 5:
        executive_summary = (
            f"The `{table_name}` table already captures live workflow data, but {row_count} rows are still too few for reliable trend conclusions. "
            "Use the current sample to validate schema quality and dashboard design, then reassess when more records accumulate."
        )
    else:
        summary_parts = [
            f"The `{table_name}` table contains {row_count} rows across {len(columns)} columns, which is enough to begin workflow-level analysis.",
        ]
        if request_columns:
            summary_parts.append(
                f"The schema is centered on request and service context fields such as {', '.join(request_columns[:3])}, which supports issue-pattern analysis and service planning."
            )
        if workflow_columns or time_columns:
            summary_parts.append(
                "The available workflow and timestamp fields can support turnaround-time, bottleneck, and SLA visibility."
            )
        if sample_truncated:
            summary_parts.append(
                "The LLM used a truncated sample, so executive decisions should be paired with SQL aggregates from the full table."
            )
        executive_summary = " ".join(summary_parts)

    growth_opportunities = _dedupe_strings([
        "Segment incoming demand by customer, sales channel, or region to identify the highest-value request sources." if customer_columns else "Create customer or segment identifiers so service demand can be analyzed by account and market.",
        "Build a recurring issue taxonomy from request details and equipment fields to identify preventive-maintenance or upsell opportunities." if request_columns else "Add structured issue-category fields so recurring demand can be grouped into commercial opportunities.",
        "Use appointment and timing data to forecast staffing needs and smooth resource allocation." if time_columns else "Add request and completion timestamps to enable demand forecasting and capacity planning.",
        "Leverage multilingual fields to support localized service delivery and cross-border operations." if language_columns else "If the business serves multiple language groups, add explicit language tracking to improve localization decisions.",
    ])

    operational_insights = _dedupe_strings([
        "Cycle-time analysis is feasible because the schema already contains time-related fields that can be converted into stage lead times." if time_columns else "Current operational visibility is limited because the table lacks enough timestamp coverage for full cycle-time analysis.",
        "The workflow can be monitored for approval or handoff bottlenecks once status or stage values are standardized." if workflow_columns else "Add explicit status or stage columns so each request can be tracked through the approval and execution pipeline.",
        f"Columns that look sparsest in the current sample are: {', '.join(sparse_columns)}." if sparse_columns else "The sampled rows look structurally consistent enough to start measuring field completeness.",
        "Because the sample sent to the model was truncated, the best practice is to pair LLM summarization with SQL-level aggregates for executive reporting." if sample_truncated else "The current sample size was small enough to inspect directly without truncation affecting interpretation.",
    ])

    customer_signals = _dedupe_strings([
        "Free-text request and condition fields can be clustered into the most common service themes and failure patterns." if request_columns else "Add narrative request-detail fields or issue tags to capture what customers are repeatedly asking for.",
        "Customer, contact, and address-related fields can reveal concentration by account, site, or geography." if customer_columns else "Customer or site identifiers are needed before concentration and repeat-demand patterns can be measured.",
        "Multilingual content suggests that language accessibility is a real service-delivery variable, not just a documentation issue." if language_columns else "If multilingual users are expected, explicit language tracking should be captured as part of the request record.",
        "The dataset does not yet appear to capture direct satisfaction or feedback, so customer sentiment is still under-observed." if not feedback_columns else "Feedback-like fields can be turned into a lightweight customer health signal once collection is made consistent.",
    ])

    risk_alerts = _dedupe_strings([
        "The dataset volume is still moderate, so it can support directional decisions but not high-confidence long-term forecasting." if row_count < 100 else "As the dataset grows, inconsistent field definitions will become a scaling risk unless taxonomy and ownership are standardized.",
        "Without explicit resolution or closure fields, management cannot reliably distinguish open backlog from completed work." if not completion_columns else "Resolution data exists, but inconsistent completion definitions would still create reporting ambiguity if not standardized.",
        "Missing priority or severity fields will make escalation and SLA breach detection harder." if not priority_columns else "Priority data should be audited for consistency before it is used in escalation dashboards.",
        "The table includes contact or location-style fields, so privacy, retention, and access control policies should be treated as first-class requirements." if customer_columns else "If customer identifiers are added later, privacy and retention controls should be designed upfront.",
    ])

    recommended_actions = _dedupe_strings([
        "Define mandatory lifecycle fields for every request: status, owner, priority, created_at, completed_at, and resolution outcome.",
        "Create a weekly operational dashboard covering request volume, cycle time, backlog, and top issue categories.",
        "Normalize free-text issues into a small controlled taxonomy so historical analysis becomes comparable over time.",
        "Review field completeness at every workflow stage and block handoff when critical fields are missing.",
    ])

    data_gaps = _dedupe_strings([
        "No explicit resolution or outcome fields detected." if not completion_columns else "Resolution fields exist, but the outcome taxonomy should still be normalized.",
        "No priority or severity tracking detected." if not priority_columns else "Priority exists, but escalation rules should be formalized around it.",
        "No direct customer feedback or satisfaction fields detected." if not feedback_columns else "Feedback fields exist, but collection quality should be checked.",
        "No obvious cost or revenue fields detected, so service economics cannot yet be evaluated." if not cost_columns else "Cost or revenue fields exist; verify that they are complete enough for unit-economics analysis.",
    ])

    dashboard_kpis = _dedupe_strings([
        "Incoming requests by day or week.",
        "Median and 90th-percentile turnaround time from request creation to completion." if time_columns else "Median turnaround time once full lifecycle timestamps are added.",
        "Backlog volume by workflow stage." if workflow_columns else "Open backlog volume after status tracking is standardized.",
        "Top issue categories and equipment or problem combinations." if request_columns else "Top standardized issue categories once taxonomy is introduced.",
    ])

    follow_up_questions = _dedupe_strings([
        "Which workflow stages matter most for management review and should become standardized status values?",
        "Which fields are mandatory before a request can move to the next approval or execution stage?",
        "Which downstream decision should this dataset support first: staffing, SLA management, recurring issue reduction, or sales follow-up?",
        "Should the business track customer satisfaction, service cost, or repeat-issue rates as the next layer of data collection?",
    ])

    return {
        "data_status": default_status,
        "executive_summary": executive_summary,
        "growth_opportunities": growth_opportunities,
        "operational_insights": operational_insights,
        "customer_signals": customer_signals,
        "risk_alerts": risk_alerts,
        "recommended_actions": recommended_actions,
        "data_gaps": data_gaps,
        "dashboard_kpis": dashboard_kpis,
        "follow_up_questions": follow_up_questions,
    }


def _merge_analysis_with_fallback(primary: Dict[str, Any], fallback: Dict[str, Any]) -> Dict[str, Any]:
    merged: Dict[str, Any] = {"data_status": primary.get("data_status") or fallback["data_status"]}
    summary = _normalize_text(primary.get("executive_summary"))
    merged["executive_summary"] = summary if len(summary) >= 30 else fallback["executive_summary"]
    for key in (
        "growth_opportunities",
        "operational_insights",
        "customer_signals",
        "risk_alerts",
        "recommended_actions",
        "data_gaps",
        "dashboard_kpis",
        "follow_up_questions",
    ):
        merged[key] = _dedupe_strings(_normalize_string_list(primary.get(key)) + fallback[key])
    return merged


def _analyze_table_report(
    table_name: str,
    row_count: int,
    columns: list[Dict[str, Any]],
    sample_rows: list[Dict[str, Any]],
    sample_truncated: bool,
    analysis_question: str,
) -> Dict[str, Any]:
    column_summary = [
        {
            "name": column["column_name"],
            "data_type": column["data_type"],
            "nullable": column["is_nullable"],
        }
        for column in columns
    ]
    prompt_payload = {
        "table_name": table_name,
        "row_count": row_count,
        "column_count": len(columns),
        "sample_rows_count": len(sample_rows),
        "sample_rows_truncated": sample_truncated,
        "analysis_question": analysis_question,
        "columns": column_summary,
        "sample_rows": sample_rows,
    }
    fallback = _fallback_analysis(table_name, row_count, columns, sample_rows, sample_truncated)
    try:
        content = _deepseek_chat(
            [
                {
                    "role": "system",
                    "content": (
                        "You are a strategic business analyst for operations, sales, product, and customer success teams. "
                        "Use only the provided table schema and sample data. "
                        "If the table is empty or sample data is limited, say so explicitly and focus on data collection priorities, "
                        "future KPIs, and practical next actions instead of inventing trends. "
                        "Return JSON only. Do not leave required fields empty."
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        "Analyze this PostgreSQL table from perspectives that are useful for future company development. "
                        "Focus on growth opportunities, customer signals, operational bottlenecks, risks, recommended actions, "
                        "data gaps, and dashboard KPIs. Return a JSON object with exactly these keys: "
                        "data_status, executive_summary, growth_opportunities, operational_insights, customer_signals, "
                        "risk_alerts, recommended_actions, data_gaps, dashboard_kpis, follow_up_questions. "
                        "`executive_summary` must be a non-empty paragraph. Each list field must contain at least 2 concise strings.\n\n"
                        f"{json.dumps(prompt_payload, ensure_ascii=False)}"
                    ),
                },
            ]
        )
        analysis = _extract_json_object(content)
    except Exception as exc:
        logger.warning("LLM table analysis fallback used for table=%s because %s", table_name, exc)
        analysis = {}

    default_status = fallback["data_status"]
    normalized_status = _normalize_key(analysis.get("data_status") or default_status) or default_status
    if normalized_status not in {"empty", "insufficient", "usable"}:
        normalized_status = default_status

    merged = _merge_analysis_with_fallback(analysis, fallback)
    merged["data_status"] = normalized_status
    return merged


def _build_analysis_markdown(
    table_name: str,
    row_count: int,
    sample_rows_count: int,
    analysis: Mapping[str, Any],
) -> str:
    sections: list[str] = [
        f"# {table_name} Analysis Report",
        "",
        f"- Row count: {row_count}",
        f"- Sample rows analyzed: {sample_rows_count}",
        f"- Data status: {analysis.get('data_status', '')}",
        "",
        "## Executive Summary",
        analysis.get("executive_summary") or "No summary returned.",
    ]
    mapping = [
        ("Growth Opportunities", analysis.get("growth_opportunities") or []),
        ("Operational Insights", analysis.get("operational_insights") or []),
        ("Customer Signals", analysis.get("customer_signals") or []),
        ("Risk Alerts", analysis.get("risk_alerts") or []),
        ("Recommended Actions", analysis.get("recommended_actions") or []),
        ("Data Gaps", analysis.get("data_gaps") or []),
        ("Dashboard KPIs", analysis.get("dashboard_kpis") or []),
        ("Follow-up Questions", analysis.get("follow_up_questions") or []),
    ]
    for title, items in mapping:
        sections.extend(["", f"## {title}"])
        if items:
            sections.extend([f"- {item}" for item in items])
        else:
            sections.append("- None")
    return "\n".join(sections)


def _build_analysis_response(
    payload: Optional[Dict[str, Any]],
    *,
    default_table_name: str,
    restrict_to_report_tables: bool,
) -> Dict[str, Any]:
    request_payload = payload or {}
    table_name = _resolve_table_name(request_payload.get("table_name") or default_table_name)
    if restrict_to_report_tables:
        table_name = _ensure_report_table_allowed(table_name)
    question = _normalize_text(request_payload.get("question")) or (
        "Focus on insights useful for company growth, customer demand, operational efficiency, risks, and next-step actions."
    )
    max_rows_raw = request_payload.get("max_rows", 30)
    try:
        max_rows = int(max_rows_raw)
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="max_rows must be an integer") from exc
    if max_rows < 1 or max_rows > 100:
        raise HTTPException(status_code=422, detail="max_rows must be between 1 and 100")

    started_at = time.time()
    with _connect_postgres() as connection:
        columns = _list_table_columns(connection, table_name)
        row_count = _count_table_rows(connection, table_name)
        raw_rows = _sample_table_rows(connection, table_name, columns, max_rows)

    prepared_rows = _prepare_rows_for_llm(raw_rows)
    analysis = _analyze_table_report(
        table_name=table_name,
        row_count=row_count,
        columns=columns,
        sample_rows=prepared_rows["rows"],
        sample_truncated=bool(prepared_rows["truncated"]),
        analysis_question=question,
    )
    report_markdown = _build_analysis_markdown(table_name, row_count, len(prepared_rows["rows"]), analysis)
    logger.info(
        "Generated info-record analysis for table=%s row_count=%s sample_rows=%s in %.2fs",
        table_name,
        row_count,
        len(prepared_rows["rows"]),
        time.time() - started_at,
    )
    return {
        "table_name": table_name,
        "row_count": row_count,
        "column_count": len(columns),
        "sample_rows_used": len(prepared_rows["rows"]),
        "sample_rows_truncated": bool(prepared_rows["truncated"]),
        "analysis_question": question,
        "analysis": analysis,
        "report_markdown": report_markdown,
        "generated_at": _utcnow_iso(),
    }


def _report_ui_index_path() -> Path:
    return REPORT_UI_DIST_DIR / "index.html"


def _serve_report_ui(path: str = "") -> FileResponse:
    index_path = _report_ui_index_path()
    if not index_path.is_file():
        raise HTTPException(status_code=503, detail="Report UI build is not available")

    requested_path = _normalize_text(path).lstrip("/")
    if requested_path:
        candidate = (REPORT_UI_DIST_DIR / requested_path).resolve()
        if REPORT_UI_DIST_DIR.resolve() in candidate.parents and candidate.is_file():
            return FileResponse(candidate)

    return FileResponse(index_path)


def _build_import_response(record_id: int) -> Dict[str, Any]:
    return {
        "record_id": record_id,
        "updated_fields": [],
        "multilingual": {},
        "warnings": [],
    }


def _render_maintenance_workbook(record: Mapping[str, Any], table_name: str) -> Dict[str, Any]:
    template_path = _get_env("MAINTENANCE_TEMPLATE_PATH", "/templates/template.xlsx")
    sheet_name = _get_env("MAINTENANCE_SHEET_NAME", "Sheet1")
    try:
        workbook = openpyxl.load_workbook(template_path)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=500, detail=f"Template file not found: {template_path}") from exc

    if sheet_name not in workbook.sheetnames:
        raise HTTPException(status_code=500, detail=f"Worksheet not found: {sheet_name}")
    worksheet = workbook[sheet_name]
    resolved_record_id = int(record["id"])
    normal_values = _build_normal_export_values(record)
    multilingual_values = {
        field_key: _prepare_multilingual_export_value(record, field_key)
        for field_key in MAINTENANCE_MULTILINGUAL_FIELD_SPECS
    }
    db_updates: Dict[str, Any] = {}
    for payload_item in multilingual_values.values():
        db_updates.update(payload_item["db_updates"])

    for field_key, value in normal_values.items():
        worksheet[MAINTENANCE_NORMAL_FIELD_SPECS[field_key]["cell"]] = value
    for payload_item in multilingual_values.values():
        worksheet[payload_item["cell"]] = payload_item["display_text"]

    metadata_text = _metadata_payload(resolved_record_id, table_name, sheet_name, normal_values, multilingual_values)
    _upsert_reply_metadata_sheet(workbook, metadata_text)

    filename = _safe_filename(table_name)
    content = _render_workbook_to_bytes(workbook)
    return {
        "record_id": resolved_record_id,
        "filename": filename,
        "content": content,
        "db_updates": db_updates,
    }


@app.post("/analyze/info-record")
def analyze_info_record(payload: Optional[Dict[str, Any]] = Body(default=None)) -> Dict[str, Any]:
    return _build_analysis_response(
        payload,
        default_table_name="info_record",
        restrict_to_report_tables=False,
    )


@app.get("/api/report/tables")
def report_table_options() -> list[Dict[str, Any]]:
    options: list[Dict[str, Any]] = []
    default_table = _report_default_table()

    with _connect_postgres() as connection:
        for table_name in _report_allowed_tables():
            try:
                columns = _list_table_columns(connection, table_name)
                row_count = _count_table_rows(connection, table_name)
            except HTTPException as exc:
                if exc.status_code == 404:
                    logger.warning("Configured report table %s was not found and will be skipped", table_name)
                    continue
                raise
            options.append(
                {
                    "name": table_name,
                    "label": _table_label(table_name),
                    "row_count": row_count,
                    "column_count": len(columns),
                    "is_default": table_name == default_table,
                }
            )

    options.sort(key=lambda item: (not item["is_default"], item["label"].lower()))
    for item in options:
        item.pop("is_default", None)
    if not options:
        raise HTTPException(status_code=500, detail="No allowed report tables are currently available")
    return options


@app.post("/api/reports/analyze")
def analyze_report(payload: Optional[Dict[str, Any]] = Body(default=None)) -> Dict[str, Any]:
    return _build_analysis_response(
        payload,
        default_table_name=_report_default_table(),
        restrict_to_report_tables=True,
    )


@app.get("/healthz")
def healthz() -> Dict[str, str]:
    return {"status": "ok"}


@app.get("/mapping/support-request")
def support_request_mapping() -> Dict[str, Any]:
    return {
        "field_cells": _load_support_request_field_map(),
        "template_path": _get_env("TEMPLATE_PATH", "/templates/Support Request Form.xlsx"),
        "sheet_name": _get_env("SHEET_NAME", "Sheet1"),
    }


@app.get("/mapping/maintenance-form")
def maintenance_form_mapping() -> Dict[str, Any]:
    return {
        "field_cells": _maintenance_field_cells(),
        "normal_fields": MAINTENANCE_NORMAL_FIELD_SPECS,
        "multilingual_fields": MAINTENANCE_MULTILINGUAL_FIELD_SPECS,
        "reply_metadata": {
            "sheet_name": REPLY_META_SHEET_NAME,
            "version": REPLY_META_VERSION,
        },
        "template_path": _get_env("MAINTENANCE_TEMPLATE_PATH", "/templates/template.xlsx"),
        "sheet_name": _get_env("MAINTENANCE_SHEET_NAME", "Sheet1"),
    }


@app.post("/fill/support-request")
def fill_support_request(payload: Dict[str, Any] = Body(...)) -> StreamingResponse:
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="Request body must be a JSON object")

    template_path = _get_env("TEMPLATE_PATH", "/templates/Support Request Form.xlsx")
    sheet_name = _get_env("SHEET_NAME", "Sheet1")
    field_map = _load_support_request_field_map()

    try:
        workbook = openpyxl.load_workbook(template_path)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=500, detail=f"Template file not found: {template_path}") from exc

    if sheet_name not in workbook.sheetnames:
        raise HTTPException(status_code=500, detail=f"Worksheet not found: {sheet_name}")

    worksheet = workbook[sheet_name]
    normalized_payload = {_normalize_key(str(key)): value for key, value in payload.items()}
    for field_key, cell in field_map.items():
        if field_key in normalized_payload:
            worksheet[cell] = _stringify(normalized_payload[field_key])

    filename_stem = _normalize_text(normalized_payload.get("customer_name")) or "support_request"
    filename = _safe_filename(f"{filename_stem}_Support_Request.xlsx")
    content = _render_workbook_to_bytes(workbook)
    output_path = _save_output_copy(content, filename)
    return _stream_xlsx(content, filename, output_path)


@app.post("/fill/maintenance-form/from-record")
def fill_maintenance_form_from_record(payload: Optional[Dict[str, Any]] = Body(default=None)) -> StreamingResponse:
    request_payload = payload or {}
    record_id_raw = request_payload.get("record_id")
    try:
        record_id = int(record_id_raw) if record_id_raw not in (None, "") else None
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="record_id must be an integer") from exc
    if not _normalize_text(request_payload.get("table_name")):
        raise HTTPException(status_code=422, detail="table_name is required")
    table_name = _resolve_table_name(request_payload.get("table_name"))

    started_at = time.time()

    with _connect_postgres() as connection:
        _ensure_maintenance_schema(connection, table_name)
        record = _fetch_record(connection, record_id, table_name)
        rendered = _render_maintenance_workbook(record, table_name)
        resolved_record_id = rendered["record_id"]
        if rendered["db_updates"]:
            _update_record_columns(connection, resolved_record_id, rendered["db_updates"], table_name)

    filename = rendered["filename"]
    content = rendered["content"]
    output_path = _save_output_copy(content, filename)
    logger.info(
        "Generated maintenance form for table=%s record_id=%s in %.2fs%s",
        table_name,
        resolved_record_id,
        time.time() - started_at,
        f" -> {output_path}" if output_path else "",
    )
    return _stream_xlsx(content, filename, output_path)


@app.post("/sync/maintenance-form/reply")
async def sync_maintenance_form_reply(
    file: UploadFile = File(...),
    sender_email: Optional[str] = Form(default=None),
    message_id: Optional[str] = Form(default=None),
) -> StreamingResponse:
    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=422, detail="Uploaded workbook is empty")

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes))
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Uploaded file is not a valid .xlsx workbook") from exc

    metadata = _load_reply_metadata(workbook)
    sheet_name = metadata.get("sheet_name")
    if not isinstance(sheet_name, str) or sheet_name not in workbook.sheetnames:
        raise HTTPException(status_code=422, detail="Workbook sheet referenced by reply metadata is missing")
    worksheet = workbook[sheet_name]

    record_id_raw = metadata.get("record_id")
    try:
        record_id = int(record_id_raw)
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="Reply metadata record_id is invalid") from exc
    if "table_name" not in metadata:
        raise HTTPException(status_code=422, detail="Reply metadata table_name is missing")
    table_name = _resolve_table_name(metadata.get("table_name"))

    normal_metadata = metadata.get("normal_fields")
    multilingual_metadata = metadata.get("multilingual_fields")
    if not isinstance(normal_metadata, dict) or not isinstance(multilingual_metadata, dict):
        raise HTTPException(status_code=422, detail="Reply metadata is incomplete")

    response_payload = _build_import_response(record_id)
    updated_field_names: set[str] = set()
    pending_updates: Dict[str, Any] = {}

    def mark_updated(name: str) -> None:
        if name not in updated_field_names:
            response_payload["updated_fields"].append(name)
            updated_field_names.add(name)

    def parse_timestamp_update(date_field_key: str, time_field_key: str, target_column: str, label: str) -> None:
        date_spec = MAINTENANCE_NORMAL_FIELD_SPECS[date_field_key]
        time_spec = MAINTENANCE_NORMAL_FIELD_SPECS[time_field_key]
        date_value = _normalize_text(worksheet[date_spec["cell"]].value)
        time_value = _normalize_text(worksheet[time_spec["cell"]].value)
        if not date_value and not time_value:
            pending_updates[target_column] = None
            mark_updated(target_column)
            return
        if not date_value or not time_value:
            raise HTTPException(
                status_code=422,
                detail=f"{label} date and time must both be filled or both be blank",
            )
        try:
            pending_updates[target_column] = _parse_form_timestamp(date_value, time_value)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=f"Invalid {label}: {exc}") from exc
        mark_updated(target_column)

    appointment_changed = {"appointment_start": False, "appointment_end": False}
    for field_key, spec in MAINTENANCE_NORMAL_FIELD_SPECS.items():
        field_meta = normal_metadata.get(field_key)
        if not isinstance(field_meta, dict):
            raise HTTPException(status_code=422, detail=f"Reply metadata missing normal field: {field_key}")
        cell = field_meta.get("cell")
        if not isinstance(cell, str) or not cell:
            raise HTTPException(status_code=422, detail=f"Reply metadata cell is invalid for {field_key}")

        current_value = _normalize_text(worksheet[cell].value)
        original_value = _normalize_text(field_meta.get("original"))
        if current_value == original_value:
            continue

        if spec["kind"] == "appointment_part":
            appointment_changed[spec["db_column"]] = True
            continue

        pending_updates[spec["db_column"]] = current_value or None
        mark_updated(field_key)

    if appointment_changed["appointment_start"]:
        parse_timestamp_update(
            "customer_appointment_date",
            "customer_appointment_time",
            "appointment_start",
            "customer appointment",
        )
    if appointment_changed["appointment_end"]:
        parse_timestamp_update("end_date", "end_time", "appointment_end", "end appointment")

    for field_key, spec in MAINTENANCE_MULTILINGUAL_FIELD_SPECS.items():
        field_meta = multilingual_metadata.get(field_key)
        if not isinstance(field_meta, dict):
            raise HTTPException(status_code=422, detail=f"Reply metadata missing multilingual field: {field_key}")
        cell = field_meta.get("cell")
        original_hashes = field_meta.get("hashes")
        if not isinstance(cell, str) or not cell:
            raise HTTPException(status_code=422, detail=f"Reply metadata cell is invalid for {field_key}")
        if not isinstance(original_hashes, dict):
            raise HTTPException(status_code=422, detail=f"Reply metadata hashes are invalid for {field_key}")

        current_value = _normalize_text(worksheet[cell].value)
        try:
            current_blocks = _parse_labeled_blocks(current_value)
        except ValueError as exc:
            raise HTTPException(
                status_code=422,
                detail=f"{field_key} must preserve the [EN], [JA], [TH] labeled structure",
            ) from exc

        changed_blocks = [
            lang
            for lang in LANGUAGE_ORDER
            if _hash_text(current_blocks.get(lang, "")) != _normalize_text(original_hashes.get(lang))
        ]
        if not changed_blocks:
            continue

        chosen_source_lang = next(lang for lang in LANGUAGE_PRIORITY if lang in changed_blocks)
        ignored_blocks = [lang for lang in changed_blocks if lang != chosen_source_lang]
        canonical_text = _normalize_text(current_blocks.get(chosen_source_lang))
        regenerated_blocks = _build_language_triplet(canonical_text, chosen_source_lang, spec["label"])

        pending_updates[spec["db_column"]] = canonical_text or None
        pending_updates[spec["source_lang_column"]] = chosen_source_lang
        for lang, column_name in spec["translation_columns"].items():
            pending_updates[column_name] = regenerated_blocks[lang]

        response_payload["multilingual"][field_key] = {
            "changed_blocks": changed_blocks,
            "chosen_source_lang": chosen_source_lang,
            "ignored_blocks": ignored_blocks,
        }
        mark_updated(field_key)

        if ignored_blocks:
            ignored_names = ", ".join(LANGUAGE_NAMES[lang].lower() for lang in ignored_blocks)
            chosen_name = LANGUAGE_NAMES[chosen_source_lang].lower()
            response_payload["warnings"].append(
                f"{field_key}: {ignored_names} edit ignored because {chosen_name} also changed"
                if len(ignored_blocks) == 1
                else f"{field_key}: {ignored_names} edits ignored because {chosen_name} also changed"
            )

    if pending_updates:
        with _connect_postgres() as connection:
            try:
                _ensure_maintenance_schema(connection, table_name)
                _update_record_columns(connection, record_id, pending_updates, table_name)
                connection.commit()
            except Exception:
                connection.rollback()
                raise

    with _connect_postgres() as connection:
        _ensure_maintenance_schema(connection, table_name)
        refreshed_record = _fetch_record(connection, record_id, table_name)
        rendered = _render_maintenance_workbook(refreshed_record, table_name)
        if rendered["db_updates"]:
            _update_record_columns(connection, record_id, rendered["db_updates"], table_name)

    if sender_email or message_id:
        logger.info(
            "Imported maintenance reply for table=%s record_id=%s sender=%s message_id=%s updated_fields=%s",
            table_name,
            record_id,
            sender_email or "",
            message_id or "",
            response_payload["updated_fields"],
        )
    else:
        logger.info(
            "Imported maintenance reply for table=%s record_id=%s updated_fields=%s",
            table_name,
            record_id,
            response_payload["updated_fields"],
        )

    output_path = _save_output_copy(rendered["content"], rendered["filename"])
    extra_headers = {
        "X-Import-Record-Id": str(record_id),
        "X-Import-Table-Name": table_name,
        "X-Updated-Fields": ",".join(response_payload["updated_fields"]),
        "X-Warnings-Count": str(len(response_payload["warnings"])),
    }
    return _stream_xlsx(rendered["content"], rendered["filename"], output_path, extra_headers)


@app.get("/", include_in_schema=False)
def report_ui_root() -> FileResponse:
    return _serve_report_ui()


@app.get("/{full_path:path}", include_in_schema=False)
def report_ui_spa(full_path: str) -> FileResponse:
    reserved_prefixes = (
        "api/",
        "analyze/",
        "fill/",
        "mapping/",
        "sync/",
        "healthz",
        "docs",
        "redoc",
        "openapi.json",
    )
    if any(full_path == prefix.rstrip("/") or full_path.startswith(prefix) for prefix in reserved_prefixes):
        raise HTTPException(status_code=404, detail="Not Found")
    return _serve_report_ui(full_path)
